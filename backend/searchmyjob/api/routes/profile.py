"""Profile operations for the workspace."""

import io
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import Response

from searchmyjob.api.dependencies import get_engine
from searchmyjob.application.ports import WorkspaceContext
from searchmyjob.domain.models import Profile

router = APIRouter()


@router.get("/api/imports/{fid}/download")
async def download_import(fid: str, *, E: WorkspaceContext = Depends(get_engine)):
    row = E.records.imported_file(fid)
    if not row:
        raise HTTPException(404)
    from urllib.parse import quote

    return Response(
        bytes(row["content"]),
        media_type=row["mime"],
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''" + quote(row["name"], safe="")
        },
    )


@router.post("/api/profile")
async def profile(body: Profile, *, E: WorkspaceContext = Depends(get_engine)):
    E.set("profile", body.text)
    return {"ok": True}


@router.post("/api/import")
async def import_cv(file: UploadFile, *, E: WorkspaceContext = Depends(get_engine)):
    data = await file.read(10_000_001)
    if len(data) > 10_000_000:
        raise HTTPException(413, "10 Mo maximum")
    suffix = Path(file.filename or "").suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages[:30])
        elif suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            text = (
                "\n".join(p.text for p in doc.paragraphs)
                + "\n"
                + "\n".join(" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
            )
        elif suffix in (".txt", ".md"):
            text = data.decode("utf-8")
        else:
            raise HTTPException(422, "Formats : PDF texte, DOCX, TXT ou Markdown")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Document illisible ou protégé")
    if not text.strip():
        raise HTTPException(422, "Aucun texte extrait ; pour un scan, colle le texte du CV")
    name = Path(file.filename or "document" + suffix).name
    mime = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
        ".md": "text/markdown",
    }[suffix]
    identifier = E.store.import_file(name, mime, data, text[:60000])
    return {"id": identifier, "text": text[:60000]}
